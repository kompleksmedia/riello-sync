#!/usr/bin/env python3
import csv
import json
import os
import posixpath
import re
import sys
import time
from decimal import Decimal, ROUND_HALF_UP
from ftplib import FTP, FTP_TLS, error_perm
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

NBP_URL = "https://api.nbp.pl/api/exchangerates/rates/A/EUR/?format=json"
CODE_RE = re.compile(r"^[A-Za-z0-9._-]+$")
WORK_DIR = Path("work")
OUTPUT_DIR = Path("output")
LOCAL_1PH = WORK_DIR / "riello_1ph.xlsx"
LOCAL_3PH = WORK_DIR / "riello_3ph.xlsx"
LOCAL_CSV = OUTPUT_DIR / "riello_prices.csv"


def env_required(name):
    value = os.getenv(name, "").strip()
    if not value:
        raise RuntimeError(f"Brak wymaganej zmiennej/sekretu: {name}")
    return value


def env_bool(name, default=False):
    raw = os.getenv(name)
    if raw is None or not raw.strip():
        return default
    return raw.strip().lower() in {"1", "true", "yes", "y", "on"}


def dec(value):
    return Decimal(str(value))


def retry(operation, attempts=3, delay=5):
    last_error = None
    for attempt in range(1, attempts + 1):
        try:
            return operation()
        except Exception as exc:
            last_error = exc
            if attempt == attempts:
                raise
            print(f"Proba {attempt}/{attempts} nieudana: {exc}. Ponawiam za {delay}s...", file=sys.stderr)
            time.sleep(delay)
    raise last_error


def connect_ftp():
    host = env_required("FTP_HOST")
    port = int(os.getenv("FTP_PORT", "21") or "21")
    user = env_required("FTP_USER")
    password = env_required("FTP_PASSWORD")
    use_tls = env_bool("FTP_USE_TLS", False)

    if use_tls:
        ftp = FTP_TLS()
        ftp.connect(host, port, timeout=30)
        ftp.login(user, password)
        ftp.prot_p()
        mode = "FTPS"
    else:
        ftp = FTP()
        ftp.connect(host, port, timeout=30)
        ftp.login(user, password)
        mode = "FTP"

    ftp.set_pasv(True)
    print(f"Polaczono z {host}:{port} przez {mode} (tryb pasywny).")
    return ftp


def download_file(remote_path, local_path):
    local_path.parent.mkdir(parents=True, exist_ok=True)

    def do_download():
        ftp = connect_ftp()
        try:
            with local_path.open("wb") as f:
                ftp.retrbinary(f"RETR {remote_path}", f.write)
        finally:
            try:
                ftp.quit()
            except Exception:
                ftp.close()

    retry(do_download)
    size = local_path.stat().st_size
    if size == 0:
        raise RuntimeError(f"Pobrany plik jest pusty: {remote_path}")
    print(f"Pobrano {remote_path} -> {local_path} ({size} B)")


def ensure_remote_directory(ftp, remote_dir):
    if not remote_dir or remote_dir in {".", "/"}:
        return

    absolute = remote_dir.startswith("/")
    parts = [p for p in remote_dir.split("/") if p]

    if absolute:
        ftp.cwd("/")

    for part in parts:
        try:
            ftp.cwd(part)
        except error_perm:
            ftp.mkd(part)
            ftp.cwd(part)


def upload_file_atomic(local_path, remote_path):
    remote_dir = posixpath.dirname(remote_path)
    remote_name = posixpath.basename(remote_path)
    if not remote_name:
        raise RuntimeError(f"Nieprawidlowa sciezka docelowa FTP: {remote_path}")
    tmp_name = remote_name + ".tmp"

    def do_upload():
        ftp = connect_ftp()
        try:
            ensure_remote_directory(ftp, remote_dir)

            try:
                ftp.delete(tmp_name)
            except error_perm:
                pass

            with local_path.open("rb") as f:
                ftp.storbinary(f"STOR {tmp_name}", f)

            try:
                ftp.rename(tmp_name, remote_name)
            except error_perm:
                try:
                    ftp.delete(remote_name)
                except error_perm:
                    pass
                ftp.rename(tmp_name, remote_name)
        finally:
            try:
                ftp.quit()
            except Exception:
                ftp.close()

    retry(do_upload)
    print(f"Wyslano {local_path} -> FTP:{remote_path}")


def get_eur_rate():
    def do_request():
        req = Request(
            NBP_URL,
            headers={
                "Accept": "application/json",
                "User-Agent": "riello-price-sync/2.0",
            },
        )
        with urlopen(req, timeout=30) as response:
            data = json.load(response)
        rate = dec(data["rates"][0]["mid"])
        effective_date = data["rates"][0].get("effectiveDate", "unknown")
        table_no = data["rates"][0].get("no", "unknown")
        return rate, effective_date, table_no

    return retry(do_request)


def find_header_positions(values):
    code_col = None
    price_col = None
    for idx, value in enumerate(values):
        if not isinstance(value, str):
            continue
        normalized = value.strip().lower()
        if normalized == "code":
            code_col = idx
        elif normalized in {"price", "priceeuro"}:
            price_col = idx
    if code_col is not None and price_col is not None:
        return code_col, price_col
    return None


def extract_prices(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    extracted = []

    try:
        for ws in wb.worksheets:
            code_col = None
            price_col = None

            for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
                header = find_header_positions(row)
                if header:
                    code_col, price_col = header
                    continue

                if code_col is None or price_col is None:
                    continue
                if max(code_col, price_col) >= len(row):
                    continue

                code = row[code_col]
                price = row[price_col]

                if not isinstance(code, str):
                    continue
                code = code.strip()
                if not code or not CODE_RE.fullmatch(code):
                    continue
                if not isinstance(price, (int, float, Decimal)):
                    continue
                if dec(price) < 0:
                    continue

                extracted.append((code, dec(price), path.name, ws.title, row_no))
    finally:
        wb.close()

    return extracted


def merge_prices(records, policy="highest"):
    merged = {}
    sources = {}
    conflicts = []

    for code, price, filename, sheet, row_no in records:
        if code not in merged:
            merged[code] = price
            sources[code] = (filename, sheet, row_no)
            continue

        old_price = merged[code]
        if old_price == price:
            continue

        conflicts.append((code, old_price, price, sources[code], (filename, sheet, row_no)))

        if policy == "highest":
            if price > old_price:
                merged[code] = price
                sources[code] = (filename, sheet, row_no)
        elif policy == "lowest":
            if price < old_price:
                merged[code] = price
                sources[code] = (filename, sheet, row_no)
        elif policy == "last":
            merged[code] = price
            sources[code] = (filename, sheet, row_no)
        elif policy == "first":
            pass
        elif policy == "error":
            raise ValueError(f"Konflikt ceny dla kodu {code}: {old_price} vs {price}")
        else:
            raise ValueError(f"Nieobslugiwana polityka duplikatow: {policy}")

    return merged, conflicts


def write_csv(output_path, merged, eur_rate):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f, delimiter=";", lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
        writer.writerow(["Code", "PriceEURO", "PricePLN"])

        for code in sorted(merged):
            euro = merged[code].quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            pln = (euro * eur_rate).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            writer.writerow([code, f"{euro:.2f}", f"{pln:.2f}"])


def main():
    path_1ph = env_required("RIELLO_1PH_FTP_PATH")
    path_3ph = env_required("RIELLO_3PH_FTP_PATH")
    output_ftp_path = env_required("RIELLO_OUTPUT_FTP_PATH")
    duplicate_policy = os.getenv("DUPLICATE_POLICY", "highest").strip().lower()

    download_file(path_1ph, LOCAL_1PH)
    download_file(path_3ph, LOCAL_3PH)

    records_1ph = extract_prices(LOCAL_1PH)
    records_3ph = extract_prices(LOCAL_3PH)
    records = records_1ph + records_3ph

    if not records:
        raise RuntimeError("Nie znaleziono zadnych poprawnych rekordow Code + Price w plikach XLSX.")

    merged, conflicts = merge_prices(records, duplicate_policy)
    eur_rate, effective_date, table_no = get_eur_rate()
    write_csv(LOCAL_CSV, merged, eur_rate)
    upload_file_atomic(LOCAL_CSV, output_ftp_path)

    print("\n=== PODSUMOWANIE ===")
    print(f"Riello 1PH - poprawne wiersze: {len(records_1ph)}")
    print(f"Riello 3PH - poprawne wiersze: {len(records_3ph)}")
    print(f"Unikalne kody w CSV: {len(merged)}")
    print(f"Kurs NBP EUR/PLN: {eur_rate}")
    print(f"Data kursu NBP: {effective_date}")
    print(f"Tabela NBP: {table_no}")
    print(f"Plik lokalny: {LOCAL_CSV}")
    print(f"Plik FTP: {output_ftp_path}")

    if conflicts:
        print(
            f"UWAGA: znaleziono {len(conflicts)} konflikt(y) cen dla powtarzajacych sie kodow. "
            f"Zastosowano polityke: {duplicate_policy}",
            file=sys.stderr,
        )
        for code, old, new, old_src, new_src in conflicts:
            print(f"  {code}: {old} {old_src} vs {new} {new_src}", file=sys.stderr)


if __name__ == "__main__":
    main()
